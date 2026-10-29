import io
from datetime import date
from decimal import Decimal

from openpyxl import load_workbook
from pypdf import PdfReader

from app.services.exports import to_csv, to_pdf, to_xlsx

COLS = [
    ("name", "Name", "text"),
    ("joined", "Joined", "date"),
    ("students", "Students", "int"),
    ("amount", "Total due", "money"),
    ("pct", "Attendance %", "percent"),
]
ROWS = [
    {
        "name": "Class 1 – A",
        "joined": date(2026, 4, 1),
        "students": 30,
        "amount": Decimal("123456.50"),
        "pct": Decimal("92.5"),
    },
    {"name": "जन्म", "joined": None, "students": 0, "amount": Decimal("0"), "pct": None},
]
TOTALS = {"students": 30, "amount": Decimal("123456.50")}


def test_csv_plain_numbers_and_utf8_bom():
    text = to_csv(COLS, ROWS).decode("utf-8")
    assert text.startswith("\ufeffName,Joined,Students,Total due,Attendance %")
    assert "Class 1 – A,2026-04-01,30,123456.50,92.5" in text  # numbers stay summable in spreadsheets
    assert "जन्म" in text


def test_xlsx_numbers_formats_and_totals():
    wb = load_workbook(io.BytesIO(to_xlsx("Class report", "Session 2026-27", COLS, ROWS, TOTALS)))
    ws = wb.active
    assert ws["A1"].value == "Class report"
    assert [c.value for c in ws[4]] == ["Name", "Joined", "Students", "Total due", "Attendance %"]
    assert ws["D5"].value == 123456.5 and ws["D5"].number_format == "#,##0.00"
    assert ws["A7"].value == "Total" and ws["D7"].value == 123456.5
    assert ws.freeze_panes == "A5"


def test_pdf_headings_values_and_totals():
    pdf = to_pdf("Green Valley", "Class report", "Session 2026-27", COLS, ROWS, TOTALS, note="first rows only")
    text = PdfReader(io.BytesIO(pdf)).pages[0].extract_text()
    for expected in ("Class report", "Total due", "₹1,23,456.50", "92.5%", "Total", "01 Apr 2026", "first rows only"):
        assert expected in text, expected


def test_formulas_in_user_text_are_not_run():
    cols = [("name", "Name", "text"), ("phone", "Phone", "text")]
    rows = [
        {"name": '=HYPERLINK("http://evil.example","x")', "phone": "+919876543210"},
        {"name": "-2+3+cmd|' /C calc'!A0", "phone": "-5"},
        {"name": "@SUM(A1)", "phone": "9876543210"},
    ]
    text = to_csv(cols, rows).decode("utf-8-sig")
    assert "'=HYPERLINK" in text and "'-2+3+cmd" in text and "'@SUM" in text
    assert "+919876543210" in text and "'+91" not in text and ",-5" in text  # plain numbers untouched
    ws = load_workbook(io.BytesIO(to_xlsx("T", "s", cols, rows, None))).active
    cells = [c for row in ws.iter_rows(min_row=5) for c in row]
    assert all(c.data_type != "f" for c in cells)
    assert ws.cell(row=5, column=1).value == '=HYPERLINK("http://evil.example","x")'  # shown as text, unchanged
