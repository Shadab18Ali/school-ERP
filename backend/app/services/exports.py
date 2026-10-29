"""CSV / Excel / PDF writers for tabular reports. Values arrive already computed
(money as Decimal, dates as date); these functions only lay them out."""

import csv
import io
import re
from datetime import UTC, date, datetime
from decimal import Decimal

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.platypus import LongTable, Paragraph, SimpleDocTemplate, Spacer, TableStyle

from app.services.receipts import _fonts, inr

Column = tuple[str, str, str]  # key, label, kind (text|int|money|percent|date)


def _cell_text(v: object, kind: str) -> str:
    if v is None:
        return ""
    if kind == "money":
        return inr(Decimal(v))
    if kind == "percent":
        return f"{v}%"
    if isinstance(v, date | datetime):
        return v.strftime("%d %b %Y")
    return str(v)


_FORMULA_START = ("=", "+", "-", "@", "\t", "\r")
_PLAIN_NUMBER = re.compile(r"^[+-]?[0-9][0-9 ,.]*$")


def safe_text(v: object) -> object:
    """Text a spreadsheet will show rather than run.

    Names and notes come from users (and from the public admission form). A cell
    starting with = + - @ is treated as a formula by Excel and LibreOffice, so it is
    prefixed with an apostrophe. Plain numbers such as +919876543210 are left alone.
    """
    if isinstance(v, str) and v.startswith(_FORMULA_START) and not _PLAIN_NUMBER.match(v):
        return "'" + v
    return v


def to_csv(columns: list[Column], rows: list[dict]) -> bytes:
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow([c[1] for c in columns])
    for r in rows:
        # plain numbers (no ₹ or grouping) so spreadsheets can sum them
        w.writerow(
            [
                "" if r.get(k) is None else (r[k].isoformat() if isinstance(r.get(k), date) else safe_text(r[k]))
                for k, _, _ in columns
            ]
        )
    return ("\ufeff" + buf.getvalue()).encode("utf-8")  # BOM: Excel opens UTF-8 (Hindi names) correctly


def to_xlsx(title: str, subtitle: str, columns: list[Column], rows: list[dict], totals: dict | None) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = title[:31]
    ws.append([title])
    ws["A1"].font = Font(bold=True, size=13)
    ws.append([subtitle])
    ws["A2"].font = Font(color="5D6B67")
    ws.append([])
    ws.append([c[1] for c in columns])
    head = 4
    for i in range(1, len(columns) + 1):
        cell = ws.cell(row=head, column=i)
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="EEF5F1")
    for r in rows:
        ws.append(
            [
                float(r[k]) if kind in ("money", "percent") and r.get(k) is not None else r.get(k)
                for k, _, kind in columns
            ]
        )
        for c in ws[ws.max_row]:
            if c.data_type == "f":  # openpyxl stores text starting with "=" as a formula: keep it text
                c.data_type = "s"
    if totals:
        ws.append(
            [
                ("Total" if i == 0 else (float(totals[k]) if k in totals and totals[k] is not None else None))
                for i, (k, _, _) in enumerate(columns)
            ]
        )
        for i in range(1, len(columns) + 1):
            ws.cell(row=ws.max_row, column=i).font = Font(bold=True)
    for i, (key, label, kind) in enumerate(columns, 1):
        letter = get_column_letter(i)
        if kind == "money":
            for c in ws[letter][head:]:
                c.number_format = "#,##0.00"
        elif kind == "percent":
            for c in ws[letter][head:]:
                c.number_format = "0.0"
        elif kind == "date":
            for c in ws[letter][head:]:
                c.number_format = "dd-mmm-yyyy"
        width = max([len(label)] + [len(_cell_text(r.get(key), kind)) for r in rows[:500]]) + 2
        ws.column_dimensions[letter].width = min(max(width, 8), 45)
        if kind in ("money", "percent", "int"):
            for c in ws[letter][head - 1 :]:
                c.alignment = Alignment(horizontal="right")
    ws.freeze_panes = ws.cell(row=head + 1, column=1)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def to_pdf(
    school: str,
    title: str,
    subtitle: str,
    columns: list[Column],
    rows: list[dict],
    totals: dict | None,
    note: str | None = None,
) -> bytes:
    regular, bold = _fonts()
    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf,
        pagesize=landscape(A4),
        leftMargin=10 * mm,
        rightMargin=10 * mm,
        topMargin=10 * mm,
        bottomMargin=10 * mm,
        title=title,
        author=school,
    )
    h = ParagraphStyle("h", fontName=bold, fontSize=13, leading=16)
    small = ParagraphStyle("s", fontName=regular, fontSize=8, leading=10, textColor=colors.HexColor("#5d6b67"))
    cell = ParagraphStyle("c", fontName=regular, fontSize=7.5, leading=9)
    cellb = ParagraphStyle("cb", fontName=bold, fontSize=7.5, leading=9)
    right = ParagraphStyle("r", parent=cell, alignment=2)
    rightb = ParagraphStyle("rb", parent=cellb, alignment=2)

    def P(v: object, kind: str, strong: bool = False, heading: bool = False) -> Paragraph:
        num = kind in ("money", "percent", "int")
        text_ = str(v) if heading else _cell_text(v, kind)  # headings are labels, never formatted as values
        return Paragraph(text_, (rightb if strong else right) if num else (cellb if strong else cell))

    data = [[P(c[1], c[2], True, heading=True) for c in columns]]
    data += [[P(r.get(k), kind) for k, _, kind in columns] for r in rows]
    if totals:
        data.append(
            [
                P("Total", "text", True) if i == 0 else P(totals.get(k), kind, True)
                for i, (k, _, kind) in enumerate(columns)
            ]
        )
    t = LongTable(data, repeatRows=1)
    style = [
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#eef5f1")),
        ("LINEBELOW", (0, 0), (-1, -1), 0.25, colors.HexColor("#dfe4e1")),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("TOPPADDING", (0, 0), (-1, -1), 2),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 2),
    ]
    if totals:
        style.append(("LINEABOVE", (0, -1), (-1, -1), 0.8, colors.HexColor("#1b362b")))
    t.setStyle(TableStyle(style))
    story = [Paragraph(school, small), Paragraph(title, h), Paragraph(subtitle, small)]
    if note:
        story.append(Paragraph(note, small))
    story += [
        Spacer(1, 4 * mm),
        t,
        Spacer(1, 3 * mm),
        Paragraph(f"{len(rows)} rows · generated {datetime.now(UTC).strftime('%d %b %Y %H:%M')} UTC", small),
    ]
    doc.build(story)
    return buf.getvalue()
